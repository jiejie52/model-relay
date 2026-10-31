"""V2 authoring workbook -> review -> approved and pinned runtime snapshot."""
from __future__ import annotations

from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.config import Settings
from app.control_plane import ModelControlPlane
from app.governance.v2_import import DeclarationError
from app.governance.v2_runtime import approve_candidate, build_candidate, digest
from app.routing.catalog import RouteCatalog


WORKBOOK = Path(__file__).resolve().parents[1] / "config" / "model_management_v2.xlsx"


def settings(**kw):
    values = dict(relay_api_token="t", supabase_url="https://example.supabase.co",
                  supabase_secret_key="t", aihubmix_api_key="t",
                  aihubmix_gemini_base_url="https://aihubmix.com/gemini")
    values.update(kw)
    return Settings(_env_file=None, **values)


def excel_bytes(change):
    wb = load_workbook(WORKBOOK)
    change(wb)
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def candidate(content=None, *, registry=None):
    return build_candidate(content or WORKBOOK, settings=settings(),
                           deployment_map={"测试": "railway", "生产": "railway-production"},
                           revision="v2-release-001", supply_registry=registry)


def approval(c):
    return {"candidate_sha256": digest(c),
            "base_control_plane_hash": c["release_metadata"]["base_control_plane_hash"],
            "approved_by": "release-operator", "approval_ref": "CHANGE-123"}


def published_settings(tmp_path, p):
    file = tmp_path / "published.json"
    file.write_text(json.dumps(p, ensure_ascii=False), encoding="utf-8")
    return settings(model_control_plane_file=str(file),
                    model_control_plane_file_sha256=sha256(file.read_bytes()).hexdigest())


def test_v2_workbook_is_configuration_source_not_test_fixture():
    assert WORKBOOK.is_file()
    assert not (Path(__file__).parent / "fixtures" / WORKBOOK.name).exists()


def test_sample_can_validate_but_remains_review_only():
    c = candidate()
    assert c["status"] == "review"
    assert c["mode"] == "merge"
    assert not c["offerings"]
    assert not c["release_metadata"]["managed_model_groups"]
    with pytest.raises(ValueError, match="status=published"):
        ModelControlPlane.from_settings(settings(model_control_plane_json=json.dumps(c)))


def test_disabled_but_published_model_shadows_builtin_wildcard(tmp_path):
    data = excel_bytes(lambda wb: wb["模型清单"].__setitem__("H6", "申请发布"))
    c = candidate(data)
    assert len(c["release_metadata"]["managed_model_groups"]) == 1
    approved = approve_candidate(c, approval(c))
    runtime = published_settings(tmp_path, approved)
    cp = ModelControlPlane.from_settings(runtime)
    routes = RouteCatalog.from_settings(runtime, control_plane=cp)
    assert routes.candidates(provider="gemini", model="gemini-3.6-flash", deployment_id="railway") == []
    assert routes.candidates(provider="gemini", model="gemini-3.8-flash", deployment_id="railway")
    # Exact environment isolation: the same model in production is not reserved.
    assert not routes._is_managed("gemini", "gemini-3.6-flash", "railway-production")
    # Both API and Worker startup paths call exactly these two constructors.
    assert ModelControlPlane.from_settings(runtime).control_plane_hash == cp.control_plane_hash


def test_legacy_full_override_cannot_eclipse_managed_group(tmp_path):
    data = excel_bytes(lambda wb: wb["模型清单"].__setitem__("H6", "申请发布"))
    c = candidate(data)
    s = published_settings(tmp_path, approve_candidate(c, approval(c)))
    s.route_catalog_json = '[{"provider":"gemini","model_pattern":"gemini-*","connection_id":"aihubmix_gemini_native"}]'
    with pytest.raises(ValueError, match="conflicts"):
        RouteCatalog.from_settings(s)


def test_wrong_hash_and_missing_snapshot_fail_closed(tmp_path):
    data = excel_bytes(lambda wb: wb["模型清单"].__setitem__("H6", "申请发布"))
    c = candidate(data)
    s = published_settings(tmp_path, approve_candidate(c, approval(c)))
    s.model_control_plane_file_sha256 = "0" * 64
    with pytest.raises(ValueError, match="SHA256 mismatch"):
        ModelControlPlane.from_settings(s)
    s.model_control_plane_file_sha256 = None
    with pytest.raises(ValueError, match="must pin"):
        ModelControlPlane.from_settings(s)
    s.model_control_plane_file = str(tmp_path / "missing.json")
    s.model_control_plane_file_sha256 = "0" * 64
    with pytest.raises(FileNotFoundError):
        ModelControlPlane.from_settings(s)


def test_file_and_legacy_json_combination_is_rejected(tmp_path):
    c = candidate()
    s = published_settings(tmp_path, approve_candidate(c, approval(c)))
    s.model_control_plane_json = '{}'
    with pytest.raises(ValueError, match="conflicts"):
        ModelControlPlane.from_settings(s)


def test_wrong_or_unapproved_manifest_rejected(tmp_path):
    c = candidate()
    with pytest.raises(DeclarationError, match="not bound"):
        approve_candidate(c, {**approval(c), "candidate_sha256": "0" * 64})
    with pytest.raises(DeclarationError, match="only review"):
        approve_candidate(approve_candidate(c, approval(c)), approval(c))
    with pytest.raises(DeclarationError, match="approver"):
        approve_candidate(c, {**approval(c), "approved_by": ""})
    # Published content cannot be injected through unpinned legacy JSON.
    release = approve_candidate(c, approval(c))
    with pytest.raises(ValueError, match="pinned"):
        ModelControlPlane.from_settings(settings(model_control_plane_json=json.dumps(release)))


def test_unknown_environment_mapping_cannot_activate():
    with pytest.raises(DeclarationError, match="no backend deployment"):
        build_candidate(WORKBOOK, settings=settings(),
                        deployment_map={"生产": "railway-production"}, revision="v2-1")


def test_enabled_route_cannot_be_published_without_certified_registry():
    def set_active(wb):
        wb["模型清单"]["H6"] = "申请发布"
        wb["渠道优先级"]["F6"] = "启用"
    with pytest.raises(DeclarationError, match="no backend-registered"):
        candidate(excel_bytes(set_active))


def test_enabled_cache_contract_cannot_silently_become_runtime_default():
    def set_active(wb):
        wb["模型清单"]["H6"] = "申请发布"
        wb["渠道优先级"]["F6"] = "启用"
        wb["缓存设置"]["I6"] = "启用"
    with pytest.raises(DeclarationError, match="cache policy"):
        candidate(excel_bytes(set_active))


def test_exact_certified_existing_supply_can_be_mapped_without_new_adapter(tmp_path):
    def set_active(wb):
        # This test activates the original first row against an already known
        # built-in Offering. Temporarily rename the separate unverified catalog
        # entry to preserve uniqueness during this synthetic certification test.
        for row in range(8, wb["模型清单"].max_row + 1):
            if wb["模型清单"].cell(row, 2).value == "gemini-3.8-flash":
                wb["模型清单"].cell(row, 2).value = "gemini-3.8-flash-catalog-placeholder"
        for name in ("渠道优先级", "缓存设置"):
            ws = wb[name]
            for row in range(10, ws.max_row + 1):
                if ws.cell(row, 1).value == "gemini-3.8-flash":
                    ws.cell(row, 1).value = "gemini-3.8-flash-catalog-placeholder"
        wb["模型清单"]["B6"] = "gemini-3.8-flash"
        wb["模型清单"]["H6"] = "申请发布"
        for row in (6, 7):
            wb["渠道优先级"][f"A{row}"] = "gemini-3.8-flash"
            wb["缓存设置"][f"A{row}"] = "gemini-3.8-flash"
        wb["渠道优先级"]["D6"] = "gemini-3.8-flash"
        wb["渠道优先级"]["F6"] = "启用"
    registry = {"schema_version": "relay-v2-supply-registry/1", "supplies": [
        {"model": "gemini-3.8-flash", "environment": "测试", "channel": "AIHubMix Gemini",
         "upstream_model": "gemini-3.8-flash", "offering_id": "aihubmix-gemini-3-8-flash",
         "certification_status": "passed", "evidence_ref": "EVIDENCE-456",
         "budget_authorization_ref": "BUDGET-789", "validated_max_output_tokens": 65536},
    ]}
    c = candidate(excel_bytes(set_active), registry=registry)
    assert len(c["offerings"]) == 1
    assert c["offerings"][0]["metadata"]["managed_v2"] is True
    cp = ModelControlPlane.from_settings(published_settings(tmp_path, approve_candidate(c, approval(c))))
    routes = RouteCatalog.from_settings(settings(), control_plane=cp)
    matching = routes.candidates(provider="gemini", model="gemini-3.8-flash", deployment_id="railway")
    assert len(matching) == 1
    assert matching[0].source == "managed_v2"
    assert matching[0].offering_id.startswith("v2:")
    assert matching[0].unavailable_action == "try_next"
    assert matching[0].capability_contract["cache"]["supported_mechanisms"] == []


def test_backend_registry_must_match_upstream_model_exactly():
    def set_active(wb):
        wb["模型清单"]["H6"] = "申请发布"
        wb["渠道优先级"]["F6"] = "启用"
    registry = {"schema_version": "relay-v2-supply-registry/1", "supplies": [
        {"model": "gemini-3.6-flash", "environment": "测试", "channel": "AIHubMix Gemini",
         "upstream_model": "gemini-3.8-flash", "offering_id": "aihubmix-gemini-3-8-flash",
         "certification_status": "passed", "evidence_ref": "EVIDENCE-456",
         "budget_authorization_ref": "BUDGET-789", "validated_max_output_tokens": 65536},
    ]}
    with pytest.raises(DeclarationError, match="alias"):
        candidate(excel_bytes(set_active), registry=registry)


def test_only_known_spreadsheet_enum_value_is_accepted():
    data = excel_bytes(lambda wb: wb["缓存设置"].__setitem__("C6", "关闭"))
    with pytest.raises(DeclarationError, match="cache mechanism"):
        candidate(data)
