"""Mainstream candidate directory is valid V2 config, never an executable supplier."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.config import Settings
from app.governance.v2_import import DeclarationError, parse_v2_declarations
from app.governance.v2_runtime import build_candidate

SOURCE = Path(__file__).resolve().parents[1] / "config" / "model_management_v2.xlsx"


def _settings():
    return Settings(_env_file=None, relay_api_token='t',
                    supabase_url='https://example.supabase.co',
                    supabase_secret_key='t', aihubmix_api_key='t',
                    aihubmix_gemini_base_url='https://aihubmix.com/gemini')


def _candidate(workbook):
    return build_candidate(workbook, settings=_settings(),
                           deployment_map={'测试':'railway','生产':'railway-production'},
                           revision='draft-catalog-only')


def test_catalog_templates_expand_table_formulas_and_data_validation_together():
    wb=load_workbook(SOURCE)
    assert wb.sheetnames == ['渠道优先级','模型清单','缓存设置','填写说明']
    for name, expected, formula in (
        ('模型清单','A5:I105','I105'),
        ('渠道优先级','A5:I105','I105'),
        ('缓存设置','A5:K105','K105')):
        ws=wb[name]
        assert next(iter(ws.tables.values())).ref == expected
        assert any('105' in str(v.sqref) for v in ws.data_validations.dataValidation)
        assert ws[formula].data_type == 'f'
        assert '$105' in ws[formula].value
        assert all('25' not in str(v.sqref) for v in ws.data_validations.dataValidation)
    assert wb['模型清单']['B6'].value == 'gemini-3.6-flash'
    assert wb['模型清单']['B7'].value == 'gemini-4-argon'
    assert 'CatalogSources' in wb['填写说明'].tables


def test_mainstream_catalog_does_not_enable_provider_routes_or_cache():
    declarations=parse_v2_declarations(SOURCE)
    assert {r['提供方'] for r in declarations['models']} >= {
        'Gemini','OpenAI','Anthropic','Qwen','DeepSeek','Kimi','GLM',
        'MiniMax','Grok','Mistral','Doubao','Baidu','Tencent','Meta Llama'}
    assert not any(r['发布意向']!='暂不发布' for r in declarations['models'])
    assert not any(r['启用意向']!='停用' for r in declarations['routes'])
    assert not any(r['启用意向']!='停用' for r in declarations['caches'])
    candidate=_candidate(SOURCE)
    assert candidate['status']=='review' and candidate['offerings']==[]
    assert candidate['release_metadata']['managed_model_groups']==[]


def test_new_provider_cannot_activate_without_exact_certified_backend_supply():
    wb=load_workbook(SOURCE)
    # Select an already-declared Qwen route; converting cells to an enable
    # intention cannot circumvent provider registration/certification.
    model_row=next(i for i in range(6,106) if wb['模型清单'].cell(i,2).value=='qwen3.8-max')
    route_row=next(i for i in range(6,106) if wb['渠道优先级'].cell(i,1).value=='qwen3.8-max')
    wb['模型清单'].cell(model_row,8).value='申请发布'
    wb['渠道优先级'].cell(route_row,6).value='启用'
    out=BytesIO(); wb.save(out)
    with pytest.raises(DeclarationError,match='no backend-registered supply'):
        _candidate(out.getvalue())
