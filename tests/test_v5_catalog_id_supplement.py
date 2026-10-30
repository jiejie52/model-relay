"""Latest screenshot and user-provided model names are draft-only, not live supplies."""
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.config import Settings
from app.governance.v2_import import DeclarationError, parse_v2_declarations
from app.governance.v2_runtime import build_candidate

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "config" / "model_management_v2.xlsx"

SCREENSHOT = {
    "gemini-3.1-flash-lite", "gemini-3.1-pro-preview",
    "gemini-3.1-pro-preview-search", "gemini-3.5-flash-lite-free",
    "gemini-3.6-flash", "grok-4.5", "grok-4.6",
    "kimi-k3", "kimi-k2.7-code",
}
TEXT_LIST = {
    "claude-opus-5-5", "claude-sonnet-5", "grok-4.7",
    "gpt-6-luna", "gpt-6-sol", "gpt-6-astra",
    "coding-glm-5.3-free", "xiaomi-mimo-v2.6-pro-free",
    "gemini-3.8-flash",
}
EXTRA = {"gpt-6.1-sol", "claude-opus-5", "mimo-v2.6-pro", "glm-5.3-flash"}
ALIASES = {"gemini-3.1-pro-preview-search", "gemini-3.5-flash-lite-free",
           "coding-glm-5.3-free", "xiaomi-mimo-v2.6-pro-free"}


def _settings():
    return Settings(_env_file=None,relay_api_token="t",
                    supabase_url="https://example.supabase.co",supabase_secret_key="t",
                    aihubmix_api_key="t",aihubmix_gemini_base_url="https://aihubmix.com/gemini")


def _candidate(source):
    return build_candidate(source,settings=_settings(),
                           deployment_map={"测试":"railway","生产":"railway-production"},
                           revision="v5-20261028-catalog-draft")


def test_user_list_and_screenshot_fully_covered_without_duplicate_ids():
    d=parse_v2_declarations(SOURCE)
    ids=[m["对外模型名"] for m in d["models"]]
    assert len(ids)==59 and len(set(ids))==59
    assert SCREENSHOT|TEXT_LIST|EXTRA<=set(ids)
    assert len(d["routes"])==61 and len(d["caches"])==61
    assert "Xiaomi" in {m["提供方"] for m in d["models"]}
    assert all(m["发布意向"]=="暂不发布" for m in d["models"])
    assert all(r["启用意向"]=="停用" for r in d["routes"])
    assert all(c["启用意向"]=="停用" for c in d["caches"])


def test_source_workbook_still_four_sheets_with_full_100_row_validation():
    wb=load_workbook(SOURCE)
    assert wb.sheetnames==["渠道优先级","模型清单","缓存设置","填写说明"]
    assert wb["填写说明"]["B33"].value=="https://z.ai/blog/glm-5.3"
    assert "ModelIdReview" in wb["填写说明"].tables
    assert wb["模型清单"]["B6"].value=="gemini-3.6-flash"
    for sheet,ref,formula in (("模型清单","A5:I105","I105"),
                               ("渠道优先级","A5:I105","I105"),
                               ("缓存设置","A5:K105","K105")):
        ws=wb[sheet]
        assert next(iter(ws.tables.values())).ref==ref
        assert all("105" in str(dv.sqref) for dv in ws.data_validations.dataValidation)
        assert ws[formula].data_type=="f"
    assert len({wb["模型清单"].cell(i,2).value for i in range(6,106) if wb["模型清单"].cell(i,2).value})==59


def test_free_search_ids_are_channel_specific_not_official_model_aliases():
    d=parse_v2_declarations(SOURCE)
    byid={m["对外模型名"]:m for m in d["models"]}
    byroute={r["对外模型名"]:r for r in d["routes"] if r["对外模型名"] in ALIASES}
    for key in ALIASES:
        assert key in byid and key in byroute
        assert byroute[key]["渠道侧模型名"]==key
        assert "聚合渠道" in byroute[key]["渠道名称"]
        assert byroute[key]["启用意向"]=="停用"
    assert byroute["xiaomi-mimo-v2.6-pro-free"]["渠道侧模型名"] != "mimo-v2.6-pro"
    assert byid["xiaomi-mimo-v2.6-pro-free"]["提供方"]=="Xiaomi"
    assert byid["mimo-v2.6-pro"]["提供方"]=="Xiaomi"
    compiled=_candidate(SOURCE)
    assert compiled["status"]=="review"
    assert compiled["offerings"]==[]
    assert compiled["release_metadata"]["managed_model_groups"]==[]


@pytest.mark.parametrize("model_id",sorted(ALIASES))
def test_unverified_aggregator_route_cannot_be_enabled_by_spreadsheet(model_id):
    wb=load_workbook(SOURCE)
    m=wb["模型清单"];r=wb["渠道优先级"]
    row=next(i for i in range(6,106) if m.cell(i,2).value==model_id)
    route=next(i for i in range(6,106) if r.cell(i,1).value==model_id)
    m.cell(row,8).value="申请发布"
    r.cell(route,6).value="启用"
    out=BytesIO();wb.save(out)
    with pytest.raises(DeclarationError,match="no backend-registered supply"):
        _candidate(out.getvalue())
