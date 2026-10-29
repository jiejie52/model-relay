"""V2 workbook declaration parser.

Produces an untrusted, non-executable draft. In particular, this module does
NOT publish a control-plane snapshot, grant capabilities, validate credentials,
or trigger inference. Server-side certification/approval is mandatory.
"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


class DeclarationError(ValueError):
    pass


_FIELDS = {
    "模型清单": ("显示名称", "对外模型名", "提供方", "输入类型", "思考档位", "结构化输出", "目标输出上限 Token", "发布意向"),
    "渠道优先级": ("对外模型名", "环境", "渠道名称", "渠道侧模型名", "顺位", "启用意向", "不可用时"),
    "缓存设置": ("对外模型名", "渠道名称", "缓存机制", "默认策略", "可选强制", "计划门槛 Token", "资源有效期 分钟", "自动准备失败时", "启用意向"),
}


def _required(value: Any, sheet: str, row: int, column: str) -> str:
    if value is None or not str(value).strip():
        raise DeclarationError(f"{sheet}!row {row}: missing {column}")
    if isinstance(value, str) and value.lstrip().startswith("="):
        raise DeclarationError(f"{sheet}!row {row}: formulas are forbidden in declaration field {column}")
    return str(value).strip()


def _positive_int(value: Any, sheet: str, row: int, column: str, *, optional: bool = False) -> int | None:
    if optional and value in (None, ""):
        return None
    if isinstance(value, bool) or not isinstance(value, int) or value <= 0:
        raise DeclarationError(f"{sheet}!row {row}: {column} must be a positive integer")
    return value


def _sheet_rows(ws, fields: tuple[str, ...]):
    # Scan the worksheet instead of relying on the original table's 20-row
    # reference; expanded declarations must receive identical validation.
    headers = [ws.cell(5, idx).value for idx in range(1, len(fields) + 1)]
    if tuple(headers) != fields:
        raise DeclarationError(f"{ws.title}: unexpected V2 header schema")
    ws.reset_dimensions()
    for idx, values in enumerate(ws.iter_rows(min_row=6, max_col=len(fields), values_only=True), start=6):
        if all(v is None or str(v).strip() == "" for v in values):
            continue
        yield idx, dict(zip(fields, values))


def parse_v2_declarations(source: bytes | str | Path, *, maximum_bytes: int = 4_000_000) -> dict[str, Any]:
    """Parse declarations without trusting Excel formulas or granting traffic.

    All output is a draft, even where a row says '启用' or '发布'. External
    publication requires a separately approved, immutable server snapshot.
    """
    payload = bytes(source) if isinstance(source, (bytes, bytearray)) else Path(source).read_bytes()
    if len(payload) > maximum_bytes:
        raise DeclarationError("workbook exceeds import size limit")
    try:
        wb = load_workbook(BytesIO(payload), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise DeclarationError("invalid V2 xlsx workbook") from exc
    try:
        if set(wb.sheetnames) != {*_FIELDS, "填写说明"}:
            raise DeclarationError("V2 workbook must have exactly four expected sheets")
        result: dict[str, Any] = {"template": "model-relay-v2", "status": "unverified_draft", "models": [], "routes": [], "caches": []}
        models: set[str] = set()
        for idx, row in _sheet_rows(wb["模型清单"], _FIELDS["模型清单"]):
            for key in ("显示名称", "对外模型名", "提供方", "输入类型", "思考档位", "结构化输出", "发布意向"):
                row[key] = _required(row[key], "模型清单", idx, key)
            row["目标输出上限 Token"] = _positive_int(row["目标输出上限 Token"], "模型清单", idx, "目标输出上限 Token", optional=True)
            key = row["对外模型名"].lower()
            if key in models:
                raise DeclarationError(f"模型清单!row {idx}: duplicate model")
            if row["发布意向"] not in ("暂不发布", "申请发布", "暂停"):
                raise DeclarationError(f"模型清单!row {idx}: invalid publication intention")

            models.add(key)
            result["models"].append(row)
        route_keys: set[tuple[str, str, str]] = set()
        ranks: set[tuple[str, str, int]] = set()
        route_pairs: set[tuple[str, str]] = set()
        for idx, row in _sheet_rows(wb["渠道优先级"], _FIELDS["渠道优先级"]):
            for key in ("对外模型名", "环境", "渠道名称", "渠道侧模型名", "启用意向", "不可用时"):
                row[key] = _required(row[key], "渠道优先级", idx, key)
            row["顺位"] = _positive_int(row["顺位"], "渠道优先级", idx, "顺位")
            model = row["对外模型名"].lower()
            if model not in models:
                raise DeclarationError(f"渠道优先级!row {idx}: model not declared")
            if row["启用意向"] not in ("启用", "停用") or row["不可用时"] not in ("尝试下一顺位", "拒绝新会话"):
                raise DeclarationError(f"渠道优先级!row {idx}: invalid enable or fallback choice")
            key = (model, row["环境"], row["渠道名称"].lower())
            rank = (model, row["环境"], row["顺位"])
            if key in route_keys or rank in ranks:
                raise DeclarationError(f"渠道优先级!row {idx}: duplicate route or priority")
            route_keys.add(key)
            route_pairs.add((model, row["渠道名称"].lower()))
            ranks.add(rank)
            result["routes"].append(row)
        cache_keys: set[tuple[str, str]] = set()
        for idx, row in _sheet_rows(wb["缓存设置"], _FIELDS["缓存设置"]):
            for key in ("对外模型名", "渠道名称", "缓存机制", "默认策略", "可选强制", "自动准备失败时", "启用意向"):
                row[key] = _required(row[key], "缓存设置", idx, key)
            model, channel = row["对外模型名"].lower(), row["渠道名称"].lower()
            key = model, channel
            if key not in route_pairs or key in cache_keys:
                raise DeclarationError(f"缓存设置!row {idx}: missing or duplicate channel declaration")
            cache_keys.add(key)
            row["计划门槛 Token"] = _positive_int(row["计划门槛 Token"], "缓存设置", idx, "计划门槛 Token", optional=True)
            row["资源有效期 分钟"] = _positive_int(row["资源有效期 分钟"], "缓存设置", idx, "资源有效期 分钟", optional=True)

            if row["缓存机制"] not in ("显式资源", "隐式前缀", "不使用"):
                raise DeclarationError(f"缓存设置!row {idx}: invalid cache mechanism")
            if row["默认策略"] not in ("关闭", "自动", "强制") or row["可选强制"] not in ("是", "否"):
                raise DeclarationError(f"缓存设置!row {idx}: invalid cache default or force flag")
            if row["启用意向"] not in ("启用", "停用"):
                raise DeclarationError(f"缓存设置!row {idx}: invalid enable choice")
            if row["自动准备失败时"] not in ("原文无缓存继续", "直接失败", "不适用"):
                raise DeclarationError(f"缓存设置!row {idx}: invalid cache prepare failure policy")
            if row["缓存机制"] == "隐式前缀" and row["资源有效期 分钟"] is not None:
                raise DeclarationError(f"缓存设置!row {idx}: implicit cache cannot have resource TTL")
            if row["默认策略"] == "强制" and (row["可选强制"] != "是" or row["缓存机制"] == "不使用"):
                raise DeclarationError(f"缓存设置!row {idx}: invalid forced cache configuration")
            result["caches"].append(row)
        return result
    finally:
        wb.close()
